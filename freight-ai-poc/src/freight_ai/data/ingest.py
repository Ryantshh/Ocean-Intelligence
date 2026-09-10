"""Explicit mappings discovered from the two TEST workbooks, not the legacy DB."""

import hashlib
import json
from collections import Counter
from datetime import datetime
from pathlib import Path

import openpyxl

from .models import Order, Provenance, Tonnage

MAPPINGS = {
    "orders": dict(
        zip(
            [
                "Date Received",
                "Update Date",
                "Lay-Can Start",
                "Lay-Can End",
                "Load / Deli",
                "Disc / Redel",
                "Cargo Types",
                "Cargo Desc.",
                "Load Parent Zone",
                "Disc Parent Zone",
                "Cargo Weight Min",
                "Cargo Weight Max",
            ],
            [
                "date_received",
                "update_date",
                "laycan_start",
                "laycan_end",
                "load_port",
                "discharge_port",
                "cargo_type",
                "cargo_description",
                "load_zone",
                "discharge_zone",
                "cargo_weight_min",
                "cargo_weight_max",
            ],
            strict=True,
        )
    ),
    "tonnage": dict(
        zip(
            [
                "Date Received",
                "First Date Received",
                "Vessel Name",
                "DWT Summer",
                "Open Areas",
                "Open Dates Start",
                "Open Dates End",
                "Update Date",
                "ETA Dates Start",
                "Parent Zone",
                "Commercial Status",
                "Ship Types",
                "Ship Sizes",
                "Status",
                "Ballast/Laden",
                "Destination",
            ],
            [
                "date_received",
                "first_date_received",
                "vessel_name",
                "dwt",
                "open_area",
                "open_date_start",
                "open_date_end",
                "update_date",
                "eta_date_start",
                "parent_zone",
                "commercial_status",
                "ship_type",
                "ship_size",
                "vessel_status",
                "ballast_laden",
                "destination",
            ],
            strict=True,
        )
    ),
}
FILENAMES = {
    "orders": "SMU_Order_data_TEST.xlsx",
    "tonnage": "SMU_Tonnage_data_TEST.xlsx",
}
SHEETS = {"orders": "Order Results", "tonnage": "Tonnage Results"}
DATE_FIELDS = {
    "laycan_start",
    "laycan_end",
    "open_date_start",
    "open_date_end",
    "eta_date_start",
}


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def read_rows(path, sheet):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=False)
    try:
        if wb.sheetnames != [sheet]:
            raise ValueError(
                f"Unexpected worksheets: {wb.sheetnames}; expected {[sheet]}"
            )
        rows = list(wb[sheet].values)
        if not rows or len(set(rows[0])) != len(rows[0]) or None in rows[0]:
            raise ValueError("Missing or duplicate header")
        for row in rows[1:]:
            if any(isinstance(v, str) and v.startswith("=") for v in row):
                raise ValueError(
                    "Formula cells require explicit review; no cached values used"
                )
        return rows[0], [
            (n, row)
            for n, row in enumerate(rows[1:], 2)
            if any(v is not None for v in row)
        ]
    finally:
        wb.close()


def inspect_sources(source_dir):
    report = {}
    for kind, name in FILENAMES.items():
        path = Path(source_dir) / name
        headers, rows = read_rows(path, SHEETS[kind])
        report[kind] = {
            "file": name,
            "sha256": digest(path),
            "sheet": SHEETS[kind],
            "rows": len(rows),
            "columns": [
                {
                    "source": h,
                    "canonical": MAPPINGS[kind].get(h),
                    "nulls": sum(row[i] is None for _, row in rows),
                    "types": dict(Counter(type(row[i]).__name__ for _, row in rows)),
                }
                for i, h in enumerate(headers)
            ],
        }
    return report


def ingest(source_dir, kind):
    path = Path(source_dir) / FILENAMES[kind]
    headers, rows = read_rows(path, SHEETS[kind])
    if set(headers) != set(MAPPINGS[kind]):
        raise ValueError(
            f"Schema drift in {path.name}: missing={set(MAPPINGS[kind]) - set(headers)}, unexpected={set(headers) - set(MAPPINGS[kind])}"
        )
    records, rejected = [], []
    sha = digest(path)
    for n, row in rows:
        values = {
            MAPPINGS[kind][h]: v.strip() if isinstance(v, str) else v
            for h, v in zip(headers, row, strict=True)
        }
        values = {k: None if v == "" else v for k, v in values.items()}
        for field in DATE_FIELDS & values.keys():
            if isinstance(values[field], datetime):
                values[field] = values[field].date()
        # Content ID identifies this snapshot; it is NOT a business order ID.
        payload = json.dumps(values, sort_keys=True, default=str)
        record_id = kind + "-" + hashlib.sha256(payload.encode()).hexdigest()[:20]
        try:
            model = Order if kind == "orders" else Tonnage
            records.append(
                model(
                    record_id=record_id,
                    source=Provenance(
                        file=path.name, sha256=sha, sheet=SHEETS[kind], row=n
                    ),
                    **values,
                )
            )
        except ValueError as exc:
            rejected.append({"file": path.name, "row": n, "error": str(exc)})
    return records, rejected


def write_json(path, value):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, indent=2, default=str) + "\n")


def preprocess(source_dir, output):
    output = Path(output)
    report = inspect_sources(source_dir)
    for kind in FILENAMES:
        records, errors = ingest(source_dir, kind)
        output.mkdir(parents=True, exist_ok=True)
        (output / f"{kind}.jsonl").write_text(
            "".join(r.model_dump_json() + "\n" for r in records)
        )
        report[kind]["processed_sha256"] = digest(output / f"{kind}.jsonl")
        report[kind]["accepted"] = len(records)
        report[kind]["rejected"] = errors
    write_json(output / "inspection.json", report)
    return report


def load_records(output, kind):
    model = Order if kind == "orders" else Tonnage
    return [
        model.model_validate_json(line)
        for line in (Path(output) / f"{kind}.jsonl").read_text().splitlines()
        if line.strip()
    ]
